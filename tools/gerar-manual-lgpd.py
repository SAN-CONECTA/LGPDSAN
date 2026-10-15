#!/usr/bin/env python3
"""Gera, para cada modelo em manual-lgpd/ (.docx e .xlsx), o .html que a tela "Manual da LGPD" abre dentro do sistema.

Uso (na raiz do repositório):  python3 tools/gerar-manual-lgpd.py
Requisitos: pandoc (https://pandoc.org) e openpyxl (pip install openpyxl).

Rode de novo sempre que trocar um .docx/.xlsx de manual-lgpd/ e suba os .html gerados junto com o arquivo original.
Os .html são fragmentos (sem <html>/<body>): o sistema os injeta na tela, como faz com biblioteca/*.html.
"""
import html
import re
import subprocess
import sys
from datetime import date, datetime
from pathlib import Path

PASTA = Path(__file__).resolve().parent.parent / 'manual-lgpd'
CAIXA = re.compile(r'<table>\s*<colgroup>\s*<col style="width: 100%" />\s*</colgroup>\s*<thead>\s*<tr>\s*<th>'
                   r'((?:(?!</th>|<table).)*)</th>\s*</tr>\s*</thead>\s*<tbody>\s*</tbody>\s*</table>', re.S)


def esc(v):
    return html.escape(str(v), quote=False).replace('\n', '<br>')


# ---------- .docx ----------
def docx_para_html(arq):
    r = subprocess.run(['pandoc', str(arq), '-f', 'docx', '-t', 'html5', '--wrap=none'],
                       capture_output=True, text=True, check=True)
    s = r.stdout
    # Os modelos usam parágrafos em negrito (não estilos de título). Promove-os a títulos para dar leitura e navegação.
    # Os dois primeiros (código e título do documento) já aparecem no cabeçalho da tela e são retirados.
    so_negrito = re.compile(r'<p><strong>([^<]+)</strong></p>\n?')
    achados = list(so_negrito.finditer(s))
    if len(achados) >= 2 and re.fullmatch(r'DOC-\d+', achados[0].group(1).strip()):
        s = s[:achados[0].start()] + s[achados[1].end():]

    def promove(m):
        t = m.group(1).strip()
        if re.match(r'^(\d+\.\s|Parte [A-Z]\b|Termo \d+\b)', t):
            return f'<h2>{t}</h2>\n'
        if re.match(r'^(Etapa \d+\b|Cláusula \d+\b|Art\. \d+)', t):
            return f'<h3>{t}</h3>\n'
        return m.group(0)
    s = so_negrito.sub(promove, s)
    # Remove classes de linha geradas pelo pandoc (odd/even/header), sem efeito aqui.
    s = re.sub(r' class="(odd|even|header)"', '', s)
    # Caixas de instrução do Word chegam como tabela de 1 coluna só com cabeçalho. Viram um aviso legível
    # (como tabela, o texto sairia todo em maiúsculas). O texto não pode cruzar </th> nem outra <table>.
    s = CAIXA.sub(lambda m: f'<div class="aviso">{m.group(1)}</div>', s)
    return s


# ---------- .xlsx ----------
def fmt(v):
    if v is None:
        return ''
    if isinstance(v, datetime):
        return v.strftime('%d/%m/%Y')
    if isinstance(v, date):
        return v.strftime('%d/%m/%Y')
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def xlsx_para_html(arq):
    import openpyxl
    wb = openpyxl.load_workbook(arq, data_only=True)
    partes = []
    for ws in wb.worksheets:
        linhas = [[fmt(c) for c in row] for row in ws.iter_rows(values_only=True)]
        cheias = [(i, l) for i, l in enumerate(linhas) if any(x.strip() for x in l)]
        # Linha de cabeçalho da tabela: a primeira com 3 ou mais células preenchidas.
        ih = next((i for i, l in cheias if sum(1 for x in l if x.strip()) >= 3), None)
        out = [f'<section class="aba" data-titulo="{html.escape(ws.title)}"><h2>{esc(ws.title)}</h2>']
        antes = [l for i, l in cheias if ih is None or i < ih]
        for k, l in enumerate(antes):
            txt = ' '.join(x.strip() for x in l if x.strip())
            out.append(f'<p><strong>{esc(txt)}</strong></p>' if k == 0 and ih is not None else f'<p>{esc(txt)}</p>')
        if ih is not None:
            cab = linhas[ih]
            ult = max(j for j, x in enumerate(cab) if x.strip())  # ignora colunas vazias à direita
            n = ult + 1
            out.append('<table class="planilha"><thead><tr>' + ''.join(f'<th>{esc(x)}</th>' for x in cab[:n]) + '</tr></thead><tbody>')
            corpo = [l for i, l in cheias if i > ih]
            for l in corpo:
                out.append('<tr>' + ''.join(f'<td>{esc(x)}</td>' for x in (l + [''] * n)[:n]) + '</tr>')
            out.append('</tbody></table>')
            em_branco = ws.max_row - (ih + 1) - len(corpo)
            if em_branco > 0:
                out.append(f'<p class="muted small">{em_branco} linhas em branco para preenchimento não aparecem aqui. '
                           'Para preencher, baixe o original (.xlsx).</p>')
        out.append('</section>')
        partes.append('\n'.join(out))
    return '\n'.join(partes) + '\n'


def main():
    if not PASTA.is_dir():
        sys.exit(f'Pasta não encontrada: {PASTA}')
    feitos = 0
    for arq in sorted(PASTA.iterdir()):
        if arq.suffix == '.docx':
            conteudo = docx_para_html(arq)
        elif arq.suffix == '.xlsx':
            conteudo = xlsx_para_html(arq)
        else:
            continue
        saida = arq.with_suffix('.html')
        saida.write_text(conteudo, encoding='utf-8')
        print(f'{arq.name} -> {saida.name} ({len(conteudo)} bytes)')
        feitos += 1
    print(f'{feitos} arquivos gerados')


if __name__ == '__main__':
    main()
